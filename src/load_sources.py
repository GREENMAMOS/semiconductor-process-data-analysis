from pathlib import Path

import openpyxl
import pandas as pd


def load_project2(path: Path) -> pd.DataFrame:
    """Return the three Si2H6 runs in tidy form."""
    try:
        wide = pd.read_excel(path, sheet_name="추가 분석 예시_Si2H6", header=1)
    except ValueError as exc:
        raise ValueError("required sheet '추가 분석 예시_Si2H6' is missing") from exc

    required = {"Process Time", "Run1", "Run2", "Run3", "Target"}
    missing = required.difference(wide.columns)
    if missing:
        raise ValueError(f"missing Project 2 columns: {sorted(missing)}")

    tidy = wide.melt(
        id_vars=["Process Time", "Target"],
        value_vars=["Run1", "Run2", "Run3"],
        var_name="run",
        value_name="si2h6_flow",
    ).rename(columns={"Process Time": "process_time", "Target": "target"})
    tidy = tidy.dropna(subset=["process_time", "si2h6_flow"])
    return tidy[["run", "process_time", "si2h6_flow", "target"]].reset_index(drop=True)


def load_project3(path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Return condition summaries and one representative wafer map."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    required_sheets = {"Sheet3", "좌표", "Sheet1"}
    missing = required_sheets.difference(wb.sheetnames)
    if missing:
        raise ValueError(f"missing Project 3 sheets: {sorted(missing)}")

    condition_rows = []
    sheet = wb["Sheet3"]
    for row in sheet.iter_rows(min_row=8, values_only=True):
        recipe = row[2] if len(row) > 2 else None
        si2h6 = row[3] if len(row) > 3 else None
        if not recipe or not isinstance(si2h6, (int, float)):
            continue
        condition_rows.append(
            {
                "recipe": recipe,
                "si2h6_flow": si2h6,
                "power": row[4],
                "temperature": row[5],
                "time_sec": row[6],
                "average_thickness": row[7],
                "max_thickness": row[8],
                "min_thickness": row[9],
                "range": row[10],
                "standard_deviation": row[11],
            }
        )

    coordinates = []
    for row in wb["좌표"].iter_rows(min_row=1, values_only=True):
        point = row[0] if len(row) > 0 else None
        x_mm = row[1] if len(row) > 1 else None
        y_mm = row[2] if len(row) > 2 else None
        if isinstance(point, (int, float)) and isinstance(x_mm, (int, float)) and isinstance(y_mm, (int, float)):
            coordinates.append((int(point), x_mm, y_mm))

    raw = wb["Sheet1"]
    thickness_row = next(
        (
            row
            for row in raw.iter_rows(min_row=2, values_only=True)
            if len(row) >= 3 and row[0] == 1 and row[1] == "THK1_1_TOP"
        ),
        None,
    )
    if thickness_row is None:
        raise ValueError("representative Project 3 thickness row is missing")

    wafer_rows = []
    for point, x_mm, y_mm in coordinates:
        value_index = point + 1
        if value_index < len(thickness_row) and isinstance(thickness_row[value_index], (int, float)):
            wafer_rows.append(
                {"point": point, "x_mm": x_mm, "y_mm": y_mm, "thickness": thickness_row[value_index]}
            )

    return pd.DataFrame(condition_rows), pd.DataFrame(wafer_rows)
