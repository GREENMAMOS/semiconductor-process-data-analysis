"""Calculations used in the detailed portfolio report (thickness in angstroms)."""
from pathlib import Path
import numpy as np
import pandas as pd
import openpyxl


def load_measurements(path):
    """Match by unique condition ID, never by repeated WF or Recipe alone.

    No implicit GOF threshold: preserve missing values and return GOF for audit.
    Input is the current educational workbook with named Korean worksheets.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        params = pd.DataFrame([
            dict(condition_id=r[1], recipe=r[2], si2h6_flow=r[3], power=r[8],
                 temperature=r[9], time_sec=r[10])
            for r in wb['공정 파라미터'].iter_rows(min_row=4, values_only=True)
            if isinstance(r[1], (int, float))
        ])
        records = []
        for r in wb['계측 데이터'].iter_rows(min_row=2, values_only=True):
            if r[2] not in ('THK1_1_TOP', 'GOF_1'):
                continue
            for p in range(1, 50):
                records.append(dict(condition_id=r[0], wf=r[1], point=p,
                                    item=r[2], value=r[p+2]))
        long = pd.DataFrame(records)
        if long.duplicated(['condition_id', 'point', 'item']).any():
            raise ValueError('Duplicate measurement key')
        thk = long[long.item == 'THK1_1_TOP'].drop(columns='item').rename(columns={'value':'thickness'})
        gof = long[long.item == 'GOF_1'][['condition_id','point','value']].rename(columns={'value':'gof'})
        joined = thk.merge(gof, on=['condition_id','point'], how='left', validate='one_to_one')
        joined = joined.merge(params, on='condition_id', how='left', validate='many_to_one', indicator=True)
        if (joined['_merge'] != 'both').any():
            raise ValueError('Unmatched condition ID')
        return joined.drop(columns='_merge')
    finally:
        wb.close()


def point_deviation(data):
    """Mean absolute deviation per location using available values, not zeros."""
    d = data.copy()
    if d.duplicated(['condition_id','point']).any():
        raise ValueError('Duplicate condition/point')
    d['wafer_mean'] = d.groupby('condition_id')['thickness'].transform('mean')
    d['absolute_deviation'] = (d.thickness-d.wafer_mean).abs()
    out = d.groupby('point').absolute_deviation.agg(
        valid_wafer_count='count', absolute_deviation_sum='sum', mean_absolute_deviation='mean'
    ).reset_index()
    return out


def evaluate_conditions(conditions, target=3200, tolerance=400):
    d = conditions.copy()
    d['uniformity_pct'] = d['range'] / (2*d.average_thickness)*100
    d['dep_rate_a_per_sec'] = d.average_thickness / d.time_sec
    d['target_gap_a'] = (d.average_thickness-target).abs()
    d['observed_points_within_spec'] = (d.min_thickness >= target-tolerance) & (d.max_thickness <= target+tolerance)
    return d


def follow_up_estimate(avg=3143.877551020407, span=623.6347590388345,
                       target=3500, tolerance=300, uniformity_drop_per_100c=.208,
                       rate_drop_per_100c=.144, base_temperature=300, base_time=45,
                       applied_temperature=234.7, applied_time=55):
    """Historical linear extrapolation, not a fitted or validated predictive model."""
    u = span/(2*avg)
    reference = tolerance/target
    drop = (1-reference/u)/uniformity_drop_per_100c*100
    expected_base_time = avg*(1-drop/100*rate_drop_per_100c)
    return dict(reference_uniformity_pct=reference*100, baseline_uniformity_pct=u*100,
                relative_uniformity_reduction_pct=(1-reference/u)*100,
                temperature_drop_c=drop, calculated_temperature_c=base_temperature-drop,
                rate_reduction_pct=drop/100*rate_drop_per_100c*100,
                expected_average_at_base_time_a=expected_base_time,
                calculated_time_sec=base_time*target/expected_base_time,
                applied_temperature_c=applied_temperature, applied_time_sec=applied_time,
                predicted_average_a=avg*applied_time/base_time*(1-(base_temperature-applied_temperature)/100*rate_drop_per_100c),
                predicted_uniformity_pct=u*(1-(base_temperature-applied_temperature)/100*uniformity_drop_per_100c)*100)
