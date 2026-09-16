from pathlib import Path

import openpyxl
import pandas as pd
import pytest

from src.load_sources import load_project2, load_project3


def test_load_project2_normalizes_three_runs(tmp_path: Path):
    path = tmp_path / "project2.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "추가 분석 예시_Si2H6"
    ws.append([None, "Si2H6"])
    ws.append(["Process Time", "Run1", "Run2", "Run3", "Target"])
    ws.append([1, 10.0, 11.0, 12.0, 10.0])
    ws.append([2, 20.0, 21.0, 22.0, 20.0])
    wb.save(path)

    result = load_project2(path)

    assert list(result.columns) == ["run", "process_time", "si2h6_flow", "target"]
    assert result.to_dict("records") == [
        {"run": "Run1", "process_time": 1, "si2h6_flow": 10.0, "target": 10.0},
        {"run": "Run1", "process_time": 2, "si2h6_flow": 20.0, "target": 20.0},
        {"run": "Run2", "process_time": 1, "si2h6_flow": 11.0, "target": 10.0},
        {"run": "Run2", "process_time": 2, "si2h6_flow": 21.0, "target": 20.0},
        {"run": "Run3", "process_time": 1, "si2h6_flow": 12.0, "target": 10.0},
        {"run": "Run3", "process_time": 2, "si2h6_flow": 22.0, "target": 20.0},
    ]


def test_load_project2_rejects_missing_sheet(tmp_path: Path):
    path = tmp_path / "bad.xlsx"
    pd.DataFrame({"x": [1]}).to_excel(path, index=False)

    with pytest.raises(ValueError, match="추가 분석 예시_Si2H6"):
        load_project2(path)


def test_load_project3_returns_conditions_and_representative_wafer(tmp_path: Path):
    path = tmp_path / "project3.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet3"
    for _ in range(6):
        ws.append([])
    ws.append([None, None, "Recipe", "Si2H6", "Power", "Temp", "Time", "AVG", "MAX", "MIN", "RANGE", "STD"])
    ws.append([None, "BKM", "R1", 200, 400, 400, 30, 100, 110, 90, 20, 5])
    ws.append([None, None, "R2", 205, 400, 400, 30, 105, 115, 95, 20, 6])
    coords = wb.create_sheet("좌표")
    coords.append([None, "x(mm)", "y(mm)"])
    coords.append([1, 0, 0])
    coords.append([2, 50, 0])
    raw = wb.create_sheet("Sheet1")
    raw.append(["WF", "ITEM_ID", 1, 2])
    raw.append([1, "THK1_1_TOP", 99, 101])
    wb.save(path)

    conditions, wafer = load_project3(path)

    assert conditions["recipe"].tolist() == ["R1", "R2"]
    assert conditions["si2h6_flow"].tolist() == [200, 205]
    assert wafer.to_dict("records") == [
        {"point": 1, "x_mm": 0, "y_mm": 0, "thickness": 99},
        {"point": 2, "x_mm": 50, "y_mm": 0, "thickness": 101},
    ]

